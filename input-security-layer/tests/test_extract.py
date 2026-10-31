import io
import json
import textwrap
from dataclasses import replace

import docx
import pymupdf
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ingest.sanitizer.extract import DocxSpan, ExtractLimits, PdfSpan, extract
from ingest.sanitizer.filetype import RejectedFile

LIMITS = ExtractLimits(
    max_file_bytes=10_000_000,
    max_pdf_pages=20,
    max_zip_uncompressed_bytes=50_000_000,
    max_text_chars=200_000,
    timeout_s=30,
    memory_mb=512,
)


# --- Fictional test files, built in memory ---

def make_pdf(pages: int = 1, white_line: bool = False, password: str | None = None) -> bytes:
    doc = pymupdf.open()
    for number in range(pages):
        page = doc.new_page()
        page.insert_text((72, 72), f"Fictional CV page {number + 1}", fontsize=11)
        if white_line:
            page.insert_text((72, 100), "hidden fictional line", fontsize=1, color=(1, 1, 1))
    if password:
        return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password, owner_pw=password)
    return doc.tobytes()


def make_docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Fictional CV: data engineer")
    hidden = document.add_paragraph().add_run("hidden fictional run")
    hidden.font.hidden = True
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Skills: Python"
    document.sections[0].header.paragraphs[0].text = "Fictional header"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def reason_for(data: bytes, file_type: str, limits: ExtractLimits = LIMITS, **kwargs) -> str:
    with pytest.raises(RejectedFile) as error:
        extract(data, file_type, limits, **kwargs)
    return error.value.reason


# --- PDF ---

def test_pdf_text_and_layout_extracted():
    spans = extract(make_pdf(white_line=True), "pdf", LIMITS)
    assert all(isinstance(span, PdfSpan) for span in spans)
    texts = [span.text for span in spans]
    assert "Fictional CV page 1" in texts
    hidden = next(span for span in spans if span.text == "hidden fictional line")
    assert hidden.color == 0xFFFFFF and hidden.size < 2  # what Step 5 will look for


def test_too_many_pages_rejected():
    assert reason_for(make_pdf(pages=3), "pdf", replace(LIMITS, max_pdf_pages=2)) == "too_many_pages"


def test_encrypted_pdf_rejected():
    assert reason_for(make_pdf(password="fictional"), "pdf") == "pdf_encrypted"


def test_corrupt_pdf_rejected():
    assert reason_for(b"%PDF-1.4\nnot really a pdf", "pdf") == "extract_failed"


# --- DOCX ---

def test_docx_runs_tables_and_header_extracted():
    spans = extract(make_docx(), "docx", LIMITS)
    assert all(isinstance(span, DocxSpan) for span in spans)
    by_text = {span.text: span for span in spans}
    assert "Fictional CV: data engineer" in by_text
    assert by_text["hidden fictional run"].hidden is True
    assert "Skills: Python" in by_text
    assert by_text["Fictional header"].part == "header"


def save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_text_hidden_through_styles_is_flagged():
    document = docx.Document()
    char_style = document.styles.add_style("FictionalHiddenChar", WD_STYLE_TYPE.CHARACTER)
    char_style.font.hidden = True
    document.add_paragraph().add_run("hidden by character style").style = char_style

    base = document.styles.add_style("FictionalHiddenBase", WD_STYLE_TYPE.PARAGRAPH)
    base.font.hidden = True
    child = document.styles.add_style("FictionalChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    document.add_paragraph("hidden by inherited paragraph style", style=child)
    document.add_paragraph("visible fictional text")

    hidden = {span.text: span.hidden for span in extract(save(document), "docx", LIMITS)}
    assert hidden["hidden by character style"] is True
    assert hidden["hidden by inherited paragraph style"] is True
    assert hidden["visible fictional text"] is False


def test_hyperlink_text_extracted():
    document = docx.Document()
    paragraph = document.add_paragraph("Email: ")
    rel_id = document.part.relate_to("mailto:amina@example.test", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "amina@example.test"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    texts = [span.text for span in extract(save(document), "docx", LIMITS)]
    assert "amina@example.test" in texts


def test_merged_cell_text_extracted_once():
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Merged fictional cell"
    texts = [span.text for span in extract(save(document), "docx", LIMITS)]
    assert texts.count("Merged fictional cell") == 1


def test_headers_read_once_and_first_page_header_included():
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Shared fictional header"
    document.sections[0].different_first_page_header_footer = True
    document.sections[0].first_page_header.paragraphs[0].text = "First page fictional header"
    document.add_section()  # its header is linked to the first one
    texts = [span.text for span in extract(save(document), "docx", LIMITS)]
    assert texts.count("Shared fictional header") == 1
    assert "First page fictional header" in texts


def test_real_decompressed_size_enforced():
    # A normal DOCX decompresses to tens of KB: over this tiny cap.
    tight = replace(LIMITS, max_zip_uncompressed_bytes=1_000)
    assert reason_for(make_docx(), "docx", tight) == "zip_bomb"


def test_too_much_text_rejected():
    assert reason_for(make_docx(), "docx", replace(LIMITS, max_text_chars=10)) == "too_much_text"


# --- The child is treated as untrusted ---

def fake_child(tmp_path, body: str):
    script = tmp_path / "fake_child.py"
    script.write_text(textwrap.dedent(body), encoding="utf-8")
    return script


def test_timeout_kills_child(tmp_path):
    child = fake_child(tmp_path, "import time; time.sleep(30)")
    limits = replace(LIMITS, timeout_s=1)
    assert reason_for(make_pdf(), "pdf", limits, child_script=child) == "extract_timeout"


def test_crashing_child_rejected(tmp_path):
    child = fake_child(tmp_path, "import sys; sys.exit(3)")
    assert reason_for(make_pdf(), "pdf", child_script=child) == "extract_failed"


def test_garbage_output_rejected(tmp_path):
    child = fake_child(tmp_path, "print('not json at all')")
    assert reason_for(make_pdf(), "pdf", child_script=child) == "extract_failed"


def test_unknown_reason_code_not_trusted(tmp_path):
    child = fake_child(tmp_path, """
        import json
        print(json.dumps({"ok": False, "reason": "<script>alert(1)</script>"}))
    """)
    assert reason_for(make_pdf(), "pdf", child_script=child) == "extract_failed"


def test_malformed_spans_rejected(tmp_path):
    child = fake_child(tmp_path, """
        import json
        print(json.dumps({"ok": True, "spans": [{"text": 123}]}))
    """)
    assert reason_for(make_pdf(), "pdf", child_script=child) == "extract_failed"


def test_child_never_sees_secrets(tmp_path, monkeypatch):
    monkeypatch.setenv("INGEST_SIGNING_KEY", "fictional-secret-" + "x" * 32)
    child = fake_child(tmp_path, """
        import json, os
        leaked = any("INGEST" in name or "IMAP" in name for name in os.environ)
        print(json.dumps({"ok": not leaked, "reason": "extract_failed", "spans": []}))
    """)
    assert extract(make_pdf(), "pdf", LIMITS, child_script=child) == []


@pytest.mark.parametrize("last_line", [
    '{"ok": false, "reason": []}',                       # unhashable reason
    "[" * 100_000 + "]" * 100_000,                        # deep nesting
    '{"ok": true, "spans": [{"text": "x", "page": 1, "line": 1, "size": NaN, "color": 0, "alpha": 255,'
    ' "char_flags": 8, "bbox": [0, 0, 1, 1], "page_box": [595, 842]}]}',  # NaN size
])
def test_hostile_child_output_rejected_cleanly(tmp_path, last_line):
    out = tmp_path / "out.txt"
    out.write_text(last_line, encoding="ascii")
    child = fake_child(tmp_path, f"import sys; sys.stdout.write(open({str(out)!r}).read())")
    assert reason_for(make_pdf(), "pdf", child_script=child) == "extract_failed"


def test_non_ascii_noise_before_result_is_ignored(tmp_path):
    child = fake_child(tmp_path, """
        import sys
        sys.stdout.buffer.write("warning: police Garamond\\u00e9\\n".encode("utf-8"))
        sys.stdout.write('{"ok": true, "spans": []}')
    """)
    assert extract(make_pdf(), "pdf", LIMITS, child_script=child) == []


def test_child_runs_in_isolated_mode(tmp_path):
    child = fake_child(tmp_path, """
        import json, sys
        print(json.dumps({"ok": bool(sys.flags.isolated), "reason": "extract_failed", "spans": []}))
    """)
    assert extract(make_pdf(), "pdf", LIMITS, child_script=child) == []
